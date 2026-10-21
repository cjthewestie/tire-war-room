"""Orders / shifts / semi-finished / operation-records database (SQLite), the 「訂單／班表資料庫」
tool of page 5. The plant has no MES (page 10), so data comes in from Excel: `import_excel()`
reads a workbook with the sheets 訂單 / 班表 / 人員 / 排程 / 半成品 / 作業紀錄, and
`export_excel()` writes the same layout (it doubles as the blank template).

Seed data reproduces the proposal's demo:
- order totals = factory.DEMO_ORDER_PROGRESS (page 2): S 38/100, M 12/100, L 8/100;
- #A-311 (small, dataset S-033) is the next mold on CURE-07, due 3 days after the demo date (page 11);
- the 胎體／胎面膠組 night shift on the demo date has 2 of 3 people (one on leave), which yields
  the gap text "10/03 夜班缺 1 人" (page 11);
- 大胎 #07 parts A 預熱中 / B 成形中 / C 包布中 / D 已完成 and goes into CURE-10 at 17:00 (page 2);
- shifts follow factory.OPTIMISED_SHIFTS (page 6 premise) or CURRENT_SHIFTS (page 4).

Everything here is read-only from the agent's point of view: the only write paths are seeding
and Excel import, both done by people.
"""
from __future__ import annotations

import sqlite3
import threading
from collections.abc import Iterable
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from warroom import factory as F
from warroom.contracts import (
    Clock,
    LaborNeed,
    Order,
    OrderProgress,
    QueryOrdersShiftsArgs,
    QueryOrdersShiftsData,
    ScheduleJob,
    SemiFinishedPart,
    SemiFinishedProgress,
    Shift,
    TeamId,
    TireSize,
)

ROOT = Path(__file__).resolve().parent.parent.parent
DEFAULT_DB_PATH = ROOT / "data" / "runtime" / "orders.db"
DEMO_DATE = date(2026, 10, 3)        # page 11 gap "10/03 夜班缺 1 人"
SIZE_LABEL = {TireSize.S: "小胎", TireSize.M: "中胎", TireSize.L: "大胎"}

SCHEMA = """
CREATE TABLE IF NOT EXISTS orders (
  id TEXT PRIMARY KEY, size TEXT NOT NULL, qty INTEGER NOT NULL, done INTEGER NOT NULL DEFAULT 0,
  due TEXT NOT NULL, customer TEXT, mold_id TEXT, priority INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS staff (
  id TEXT PRIMARY KEY, name TEXT NOT NULL, team TEXT NOT NULL, shift_name TEXT NOT NULL,
  can_demold INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS shifts (
  id TEXT PRIMARY KEY, team TEXT NOT NULL, name TEXT NOT NULL, start TEXT NOT NULL, end TEXT NOT NULL,
  headcount INTEGER NOT NULL, required INTEGER);
CREATE TABLE IF NOT EXISTS absences (
  staff_id TEXT NOT NULL, shift_id TEXT NOT NULL, reason TEXT, PRIMARY KEY (staff_id, shift_id));
CREATE TABLE IF NOT EXISTS schedule_jobs (
  id TEXT PRIMARY KEY, order_id TEXT NOT NULL, size TEXT NOT NULL, press_id TEXT NOT NULL,
  mold_id TEXT, mold_in TEXT NOT NULL, cure_h REAL NOT NULL, frozen INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS semi_finished (
  tire_label TEXT NOT NULL, order_id TEXT NOT NULL, size TEXT NOT NULL, part TEXT NOT NULL,
  label TEXT NOT NULL, status_text TEXT NOT NULL, state TEXT NOT NULL, machine_id TEXT,
  planned_mold_in_at TEXT, PRIMARY KEY (tire_label, part));
CREATE TABLE IF NOT EXISTS operation_records (
  id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, machine_id TEXT, kind TEXT NOT NULL,
  team TEXT, operator TEXT, order_id TEXT, note TEXT);
"""

# Excel layout: sheet name -> (table, columns). Header row uses these column names.
EXCEL_SHEETS: dict[str, tuple[str, tuple[str, ...]]] = {
    "訂單": ("orders", ("id", "size", "qty", "done", "due", "customer", "mold_id", "priority")),
    "人員": ("staff", ("id", "name", "team", "shift_name", "can_demold")),
    "班表": ("shifts", ("id", "team", "name", "start", "end", "headcount", "required")),
    "請假": ("absences", ("staff_id", "shift_id", "reason")),
    "排程": ("schedule_jobs", ("id", "order_id", "size", "press_id", "mold_id", "mold_in", "cure_h", "frozen")),
    "半成品": ("semi_finished", ("tire_label", "order_id", "size", "part", "label", "status_text", "state",
                               "machine_id", "planned_mold_in_at")),
    "作業紀錄": ("operation_records", ("ts", "machine_id", "kind", "team", "operator", "order_id", "note")),
}
# Accept Chinese header aliases from a plant's own spreadsheet.
HEADER_ALIASES = {
    "訂單號": "id", "訂單編號": "id", "尺寸": "size", "數量": "qty", "已完成": "done", "交期": "due",
    "客戶": "customer", "模具": "mold_id", "模號": "mold_id", "優先": "priority",
    "工號": "id", "姓名": "name", "組別": "team", "班別": "shift_name", "可出入模": "can_demold",
    "班次編號": "id", "名稱": "name", "開始": "start", "結束": "end", "人數": "headcount", "應到": "required",
    "機台": "press_id", "入模": "mold_in", "硫化時數": "cure_h", "時間": "ts", "類別": "kind",
    "作業員": "operator", "備註": "note",
}


def _iso(v: Any) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return (v if v.tzinfo else v.replace(tzinfo=F.TZ)).isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    try:
        d = datetime.fromisoformat(s)
        return (d if d.tzinfo else d.replace(tzinfo=F.TZ)).isoformat()
    except ValueError:
        return s


def _dt(s: str) -> datetime:
    d = datetime.fromisoformat(s)
    return d.astimezone(F.TZ) if d.tzinfo else d.replace(tzinfo=F.TZ)


def _date(s: Any) -> date:
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    return date.fromisoformat(str(s)[:10])


def build_job(job_id: str, order_id: str, size: TireSize, press_id: str, mold_in: datetime,
              cure_h: float | None = None, mold_id: str | None = None, frozen: bool = False) -> ScheduleJob:
    """ScheduleJob with page-6 back-propagation. Uses warroom.schedule.back_propagate when that
    module is implemented, otherwise the same table from factory.BACKPROP."""
    cure = cure_h if cure_h is not None else F.CURE_HOURS[size]
    try:
        from warroom.schedule import back_propagate
        job = back_propagate(job_id, order_id, size, press_id, mold_in, cure_h=cure, mold_id=mold_id)
        return job.model_copy(update={"frozen": frozen})
    except (ImportError, NotImplementedError, AttributeError):
        pass
    bp = F.BACKPROP[size]
    bond = mold_in - timedelta(minutes=bp["bond_start_min"])
    feed = mold_in - timedelta(minutes=bp["a_feed_min"])
    demold = mold_in + timedelta(hours=cure)
    labor = [
        LaborNeed(team=TeamId.BODY, task="a_feed", start=feed, end=bond),
        LaborNeed(team=TeamId.BOND, task="bonding", start=bond, end=mold_in),
        LaborNeed(team=_team_on_duty(mold_in), task="mold_in", start=mold_in,
                  end=mold_in + timedelta(minutes=F.DEMOLD_MOLDIN_MIN)),
        LaborNeed(team=_team_on_duty(demold), task="demold", start=demold,
                  end=demold + timedelta(minutes=F.DEMOLD_MOLDIN_MIN)),
    ]
    return ScheduleJob(id=job_id, order_id=order_id, size=size, press_id=press_id, mold_id=mold_id,
                       mold_in=mold_in, demold=demold, cure_h=cure, bond_start=bond, a_feed_start=feed,
                       side_pieces=bp["side_pieces"], bead_pieces=bp["bead_pieces"], labor=labor, frozen=frozen)


def _team_on_duty(t: datetime) -> TeamId:
    """Curing team 06–23 (staggered), 胎體／胎面膠組 night 20–05 does de-mold / mold-in (page 6)."""
    h = t.astimezone(F.TZ).hour
    return TeamId.CURING if 6 <= h < 23 else TeamId.BODY


class OrdersDB:
    """SQLite store. path=":memory:" for tests. Thread-safe (one connection, one lock)."""

    def __init__(self, path: Path | str | None = None, *, ref_date: date | None = None,
                 clock: Clock | None = None, seed: bool = True, shift_plan: str = "optimised",
                 xlsx: Path | None = None):
        self.path = str(path if path is not None else DEFAULT_DB_PATH)
        if self.path != ":memory:":
            Path(self.path).parent.mkdir(parents=True, exist_ok=True)
        self.clock = clock
        self.ref_date = ref_date or (clock.now().astimezone(F.TZ).date() if clock else DEMO_DATE)
        self._lock = threading.RLock()
        self.conn = sqlite3.connect(self.path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        if xlsx:
            self.import_excel(xlsx)
        elif seed and self._count("orders") == 0:
            self.seed(self.ref_date, shift_plan=shift_plan)

    # ------------------------------------------------------------------ basics
    def _count(self, table: str) -> int:
        return self.conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

    def _rows(self, sql: str, params: Iterable[Any] = ()) -> list[sqlite3.Row]:
        with self._lock:
            return self.conn.execute(sql, tuple(params)).fetchall()

    def clear(self) -> None:
        with self._lock:
            for table, _ in EXCEL_SHEETS.values():
                self.conn.execute(f"DELETE FROM {table}")
            self.conn.commit()

    def now(self) -> datetime:
        if self.clock:
            return self.clock.now()
        return datetime.combine(self.ref_date, time(0), tzinfo=F.TZ)

    # ------------------------------------------------------------------ seed
    def seed(self, ref: date | None = None, shift_plan: str = "optimised") -> None:
        """Demo data anchored on ref (default 2026-10-03, the page-11 night)."""
        ref = ref or self.ref_date
        self.ref_date = ref
        self.clear()
        D = lambda d=0, h=0, m=0: datetime.combine(ref + timedelta(days=d), time(h, m), tzinfo=F.TZ)
        orders = [
            # small 38/100
            ("#S-101", "S", 40, 30, ref + timedelta(days=4), "興旺堆高機股份有限公司", "S-24A", 1),
            ("#S-102", "S", 30, 8, ref + timedelta(days=9), "阿財搬運車行", "S-24B", 0),
            ("#S-103", "S", 26, 0, ref + timedelta(days=16), "順手物流設備有限公司", "S-24A", 0),
            ("#A-311", "S", 4, 0, ref + timedelta(days=3), "永固機械股份有限公司", "S-24A", 3),   # next mold on CURE-07, due in 3 days (page 11)
            # medium 12/100
            ("#A-309", "M", 20, 12, ref + timedelta(days=12), "港都機具工業有限公司", "M-28A", 1),
            ("#A-315", "M", 40, 0, ref + timedelta(days=40), "元氣滿滿物流股份有限公司", "M-28B", 0),
            ("#A-318", "M", 40, 0, ref + timedelta(days=60), "海派貨櫃場有限公司", "M-28B", 0),
            # large 8/100: 大胎 #07 belongs to #B-207 (page 2)
            ("#B-207", "L", 50, 8, ref + timedelta(days=45), "青海碼頭作業有限公司", "L-34A", 2),
            ("#B-212", "L", 50, 0, ref + timedelta(days=80), "鋼鐵人運輸股份有限公司", "L-34A", 0),
        ]
        self._insert("orders", EXCEL_SHEETS["訂單"][1],
                     [(o[0], o[1], o[2], o[3], o[4].isoformat(), o[5], o[6], o[7]) for o in orders])

        shift_defs = F.OPTIMISED_SHIFTS if shift_plan == "optimised" else F.CURRENT_SHIFTS
        staff, shifts, absences = [], [], []
        names = iter(_NAMES)
        for sd in shift_defs:
            n = 3 if (sd.team is TeamId.BODY and sd.end < sd.start) else sd.headcount   # night does de-mold
            members = []
            for i in range(n):
                sid = f"{sd.team.value[:4].upper()}-{_SHIFT_CODE.get(sd.name, 'X')}{i + 1}"
                members.append(sid)
                staff.append((sid, next(names), sd.team.value, sd.name,
                              int(sd.team in (TeamId.CURING, TeamId.BODY))))
            for d in range(-1, 7):
                day = ref + timedelta(days=d)
                if day.weekday() not in sd.weekdays:
                    continue
                start = datetime.combine(day, sd.start, tzinfo=F.TZ)
                end = datetime.combine(day + timedelta(days=1 if sd.end <= sd.start else 0), sd.end, tzinfo=F.TZ)
                shift_id = f"{day:%m%d}-{sd.team.value}-{sd.name}"
                headcount = n
                if d == 0 and sd.team is TeamId.BODY and sd.end < sd.start:
                    absences.append((members[-1], shift_id, "請假（病假）"))   # page 11: 夜班缺 1 人
                    headcount = n - 1
                shifts.append((shift_id, sd.team.value, sd.name, start.isoformat(), end.isoformat(), headcount, n))
        self._insert("staff", EXCEL_SHEETS["人員"][1], staff)
        self._insert("shifts", EXCEL_SHEETS["班表"][1], shifts)
        self._insert("absences", EXCEL_SHEETS["請假"][1], absences)

        P = F.press_id
        jobs = [
            # CURE-07 (small-only, dataset EV-2 S-027): small #S-101 at 128 °C, mold-in 11:50 the day before
            # -> 14.4 h of 24 h in at 02:14 (page 11); #A-311 is its next mold
            ("J-07-01", "#S-101", "S", P(7), "S-24A", D(-1, 11, 50), 24.0, 1),
            ("J-07-02", "#A-311", "S", P(7), "S-24A", D(0, 20, 44), 24.0, 0),
            ("J-11-01", "#A-309", "M", P(11), "M-28A", D(-1, 16, 14), 28.0, 1),
            ("J-11-02", "#A-315", "M", P(11), "M-28B", D(0, 20, 44), 28.0, 0),
            ("J-08-01", "#A-309", "M", P(8), "M-28A", D(-1, 8, 30), 28.0, 1),
            ("J-08-02", "#A-315", "M", P(8), "M-28B", D(0, 13, 0), 28.0, 0),
            ("J-09-01", "#B-207", "L", P(9), "L-34A", D(-1, 14, 0), 34.0, 1),
            ("J-09-02", "#B-207", "L", P(9), "L-34A", D(1, 0, 30), 34.0, 0),
            # CURE-10: 大胎 #07 goes in at 17:00 (page 2 "大胎 #07 仍 17:00 入硫化機")
            ("J-10-01", "#B-207", "L", P(10), "L-34A", D(-1, 6, 30), 34.0, 1),
            ("J-10-07", "#B-207", "L", P(10), "L-34A", D(0, 17, 0), 34.0, 0),
        ]
        for i, n in enumerate((1, 2, 3, 4, 5, 6)):
            order = "#S-101" if i < 4 else "#S-102"
            start = D(-1, 7 + i, 0)
            jobs.append((f"J-{n:02d}-01", order, "S", P(n), "S-24A" if i < 4 else "S-24B", start, 24.0, 1))
            jobs.append((f"J-{n:02d}-02", order, "S", P(n), "S-24A" if i < 4 else "S-24B",
                         start + timedelta(hours=24, minutes=30), 24.0, 0))
        self._insert("schedule_jobs", EXCEL_SHEETS["排程"][1],
                     [(j[0], j[1], j[2], j[3], j[4], j[5].isoformat(), j[6], j[7]) for j in jobs])

        mold_in_07 = D(0, 17, 0).isoformat()
        semi = [
            ("大胎 #07", "#B-207", "L", "A", F.LINE_LABELS[F.Line.A], "預熱中", "in_progress", "A-BIGOPEN", mold_in_07),
            ("大胎 #07", "#B-207", "L", "B", F.LINE_LABELS[F.Line.B], "成形中", "in_progress", "B-FORM", mold_in_07),
            ("大胎 #07", "#B-207", "L", "C", F.LINE_LABELS[F.Line.C], "包布中", "in_progress", "C-WRAP", mold_in_07),
            ("大胎 #07", "#B-207", "L", "D", F.LINE_LABELS[F.Line.D], "已完成", "done", "D-CUT", mold_in_07),
        ]
        self._insert("semi_finished", EXCEL_SHEETS["半成品"][1], semi)

        records = [
            (D(-1, 6, 30), P(10), "入模", "curing", "CURI-早1", "#B-207", "大胎入模，設定 128°C"),
            (D(-1, 8, 30), P(8), "入模", "curing", "CURI-早2", "#A-309", "中胎入模"),
            (D(-1, 10, 5), "A-BIGOPEN", "巡檢", "bigopen", "BIGO-早1", None, "膠溫 96°C，水流正常"),
            (D(-1, 14, 0), P(9), "入模", "curing", "CURI-中1", "#B-207", "大胎入模"),
            (D(-1, 15, 40), "A-MIX", "QC 取樣", "qc", None, "#B-207", "門尼黏度 62，合格"),
            (D(-1, 11, 50), P(7), "入模", "curing", "CURI-中2", "#S-101", "小胎入模，128°C，預計 24 h"),
            (D(-1, 18, 0), P(3), "保養", "maintenance", None, None, "更換疏水閥墊片"),
            (D(-1, 21, 30), P(6), "巡檢", "body", "BODY-夜1", None, "6 號機 128°C，蒸汽壓正常"),
            (D(0, 0, 10), "A-SMALLOPEN", "巡檢", "body", "BODY-夜2", None, "保溫儲料 85°C"),
        ]
        self._insert("operation_records", EXCEL_SHEETS["作業紀錄"][1],
                     [(r[0].isoformat(), *r[1:]) for r in records])

    def _insert(self, table: str, cols: tuple[str, ...], rows: list[tuple]) -> int:
        if not rows:
            return 0
        sql = f"INSERT OR REPLACE INTO {table} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})"
        with self._lock:
            self.conn.executemany(sql, rows)
            self.conn.commit()
        return len(rows)

    # ------------------------------------------------------------------ Excel
    def import_excel(self, path: Path, replace: bool = True) -> dict[str, int]:
        """Import a workbook (no MES on site). Missing sheets are left untouched. Returns rows per sheet."""
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        counts: dict[str, int] = {}
        for sheet, (table, cols) in EXCEL_SHEETS.items():
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            if not rows:
                continue
            header = [HEADER_ALIASES.get(str(h).strip(), str(h).strip()) if h is not None else "" for h in rows[0]]
            out = []
            for r in rows[1:]:
                if r is None or all(v in (None, "") for v in r):
                    continue
                rec = dict(zip(header, r))
                vals = []
                for c in cols:
                    v = rec.get(c)
                    if c in ("due",):
                        v = _date(v).isoformat()
                    elif c in ("start", "end", "mold_in", "ts", "planned_mold_in_at") or isinstance(v, datetime):
                        v = _iso(v)
                    vals.append(v)
                out.append(tuple(vals))
            if replace:
                with self._lock:
                    self.conn.execute(f"DELETE FROM {table}")
            counts[sheet] = self._insert(table, cols, out)
        wb.close()
        return counts

    def export_excel(self, path: Path) -> Path:
        """Write every table to a workbook in the import layout (also the blank template)."""
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        for sheet, (table, cols) in EXCEL_SHEETS.items():
            ws = wb.create_sheet(sheet)
            ws.append(list(cols))
            for r in self._rows(f"SELECT {','.join(cols)} FROM {table}"):
                ws.append(list(r))
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return Path(path)

    # ------------------------------------------------------------------ reads
    def orders(self, order_ids: list[str] | None = None, open_only: bool = False) -> list[Order]:
        rows = self._rows("SELECT * FROM orders ORDER BY due, id")
        out = [Order(id=r["id"], size=TireSize(r["size"]), qty=r["qty"], done=r["done"], due=_date(r["due"]),
                     customer=r["customer"], mold_id=r["mold_id"], priority=r["priority"] or 0) for r in rows]
        if order_ids:
            want = {_norm_order(o) for o in order_ids}
            out = [o for o in out if _norm_order(o.id) in want]
        if open_only:
            out = [o for o in out if o.done < o.qty]
        return out

    def shifts(self, from_ts: datetime | None = None, to_ts: datetime | None = None,
               teams: list[TeamId] | None = None) -> list[Shift]:
        out = []
        for r in self._rows("SELECT * FROM shifts ORDER BY start, team"):
            s = Shift(id=r["id"], team=TeamId(r["team"]), name=r["name"], start=_dt(r["start"]),
                      end=_dt(r["end"]), headcount=r["headcount"], required=r["required"])
            if from_ts and s.end <= from_ts:
                continue
            if to_ts and s.start >= to_ts:
                continue
            if teams and s.team not in teams:
                continue
            out.append(s)
        return out

    def staff(self, team: TeamId | None = None) -> list[dict[str, Any]]:
        rows = self._rows("SELECT * FROM staff ORDER BY team, shift_name, id")
        return [dict(r) for r in rows if team is None or r["team"] == team.value]

    def absences(self) -> list[dict[str, Any]]:
        return [dict(r) for r in self._rows("SELECT * FROM absences ORDER BY shift_id")]

    def jobs(self, press_ids: list[str] | None = None, order_ids: list[str] | None = None,
             after: datetime | None = None) -> list[ScheduleJob]:
        out = []
        for r in self._rows("SELECT * FROM schedule_jobs ORDER BY mold_in, press_id"):
            if press_ids and r["press_id"] not in press_ids:
                continue
            if order_ids and _norm_order(r["order_id"]) not in {_norm_order(o) for o in order_ids}:
                continue
            job = build_job(r["id"], r["order_id"], TireSize(r["size"]), r["press_id"], _dt(r["mold_in"]),
                            cure_h=r["cure_h"], mold_id=r["mold_id"], frozen=bool(r["frozen"]))
            if after and job.demold <= after:
                continue
            out.append(job)
        return out

    def gaps(self, shifts: list[Shift]) -> list[str]:
        """["10/03 夜班缺 1 人"] for every shift staffed below its required headcount."""
        out = []
        for s in shifts:
            if s.required is not None and s.headcount < s.required:
                out.append(f"{s.start:%m/%d} {s.name}缺 {s.required - s.headcount} 人")
        return out

    def gap_details(self, shifts: list[Shift]) -> list[dict[str, Any]]:
        return [{"shift_id": s.id, "date": s.start.date().isoformat(), "shift_name": s.name,
                 "team": s.team.value, "team_label": F.TEAM_LABELS[s.team],
                 "headcount": s.headcount, "required": s.required, "short": s.required - s.headcount}
                for s in shifts if s.required is not None and s.headcount < s.required]

    def query(self, args: QueryOrdersShiftsArgs) -> QueryOrdersShiftsData:
        """The query_orders_shifts tool (R0, read-only). Default window: now → +48 h."""
        now = self.now()
        frm = args.from_ts or now
        to = args.to_ts or (frm + timedelta(hours=48))
        if args.press_ids:
            jobs = self.jobs(press_ids=args.press_ids, after=now)
            jobs = _next_per_press(jobs, now)
        elif args.order_ids:
            jobs = self.jobs(order_ids=args.order_ids, after=now)
        else:
            jobs = self.jobs(after=now)[:12]
        if args.order_ids:
            orders = self.orders(args.order_ids)
        elif args.press_ids:
            ids = {j.order_id for j in jobs}
            orders = [o for o in self.orders() if o.id in ids]
        else:
            orders = self.orders(open_only=True)
        shifts = self.shifts(frm, to)
        return QueryOrdersShiftsData(orders=orders, shifts=shifts, gaps=self.gaps(shifts), next_jobs=jobs)

    def semi_finished(self, tire_label: str | None = None) -> SemiFinishedProgress | None:
        rows = self._rows("SELECT * FROM semi_finished ORDER BY tire_label, part")
        if tire_label:
            rows = [r for r in rows if r["tire_label"] == tire_label]
        if not rows:
            return None
        first = rows[0]
        rows = [r for r in rows if r["tire_label"] == first["tire_label"]]
        return SemiFinishedProgress(
            tire_label=first["tire_label"], order_id=first["order_id"], size=TireSize(first["size"]),
            parts=[SemiFinishedPart(part=r["part"], label=r["label"], status_text=r["status_text"],
                                    state=r["state"], machine_id=r["machine_id"]) for r in rows],
            planned_mold_in_at=_dt(first["planned_mold_in_at"]) if first["planned_mold_in_at"] else None)

    def order_progress(self) -> list[OrderProgress]:
        """Page 2 panel ①. remaining_h is the proposal's figure when the seeded totals match it,
        otherwise an estimate from page-6 weekly capacity (AI 出模時刻 row)."""
        target = {r.label: r for r in F.CAPACITY_TARGETS}["AI 排出模時刻"]
        per_press_week = {TireSize.S: target.small, TireSize.M: target.medium, TireSize.L: target.large}
        presses = {TireSize.S: len(F.SMALL_ONLY_PRESSES), TireSize.M: 2, TireSize.L: 2}
        out = []
        for size in TireSize:
            os_ = [o for o in self.orders() if o.size is size]
            done, total = sum(o.done for o in os_), sum(o.qty for o in os_)
            demo = F.DEMO_ORDER_PROGRESS[size]
            if (done, total) == demo[:2]:
                rem = float(demo[2])
            else:
                rem = round((total - done) / (per_press_week[size] * presses[size]) * 168.0, 1)
            out.append(OrderProgress(size=size, label=SIZE_LABEL[size], done=done, total=total, remaining_h=rem))
        return out

    def operation_records(self, machine_id: str | None = None, since: datetime | None = None,
                          limit: int = 50) -> list[dict[str, Any]]:
        rows = self._rows("SELECT * FROM operation_records ORDER BY ts DESC, id DESC")
        out = []
        for r in rows:
            if machine_id and r["machine_id"] != machine_id:
                continue
            if since and _dt(r["ts"]) < since:
                continue
            out.append(dict(r))
            if len(out) >= limit:
                break
        return out


def _norm_order(o: str) -> str:
    return o.strip().lstrip("#").upper()


def _next_per_press(jobs: list[ScheduleJob], now: datetime) -> list[ScheduleJob]:
    """Running job + next upcoming job per press (what the commander needs to reschedule)."""
    out: list[ScheduleJob] = []
    seen: dict[str, int] = {}
    for j in jobs:
        n = seen.get(j.press_id, 0)
        if n < 2:
            out.append(j)
            seen[j.press_id] = n + 1
    return out


_SHIFT_CODE = {"早班": "早", "中班": "中", "夜班": "夜", "晚班": "晚"}
_NAMES = ("陳志明", "林俊傑", "黃建宏", "張家豪", "李宗翰", "王志偉", "吳承恩", "劉冠廷", "蔡明哲", "楊凱文",
          "許文彬", "鄭宇軒", "謝佳穎", "郭俊宏", "洪嘉文", "曾國峰", "邱文志", "廖偉傑", "賴怡君", "周世豪",
          "徐子軒", "蘇育賢", "葉建良", "莊雅婷", "呂宗憲", "江柏翰", "何建志", "高士凱", "羅偉倫", "鍾雅雯")
